"""Reopen a generated deck and report structure/content for verification."""
import sys

from pptx import Presentation


def main():
    path = sys.argv[1]
    prs = Presentation(path)
    print(f"file: {path}")
    print(f"slide size: {prs.slide_width/914400:.2f} x {prs.slide_height/914400:.2f} in")
    print(f"slides: {len(prs.slides)}")
    print()
    problems = 0
    for i, s in enumerate(prs.slides, 1):
        title = ""
        if s.shapes.title is not None and s.shapes.title.text_frame.text.strip():
            title = s.shapes.title.text_frame.text.strip().replace("\n", " / ")
        texts = []
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    t = p.text.strip()
                    if t and t != title:
                        texts.append(t)
        # overflow-ish check: any shape extending past the slide
        for sh in s.shapes:
            try:
                if sh.left is not None and sh.left + sh.width > prs.slide_width + 10000:
                    print(f"  !! slide {i}: shape {sh.name!r} runs past right edge")
                    problems += 1
                if sh.top is not None and sh.top + sh.height > prs.slide_height + 10000:
                    print(f"  !! slide {i}: shape {sh.name!r} runs past bottom edge")
                    problems += 1
            except TypeError:
                pass
        print(f"  {i:2d}. [{s.slide_layout.name}] {title[:56]}")
        print(f"      blocks={len(texts)} chars={sum(len(t) for t in texts)}")
    print()
    print(f"geometry problems: {problems}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
